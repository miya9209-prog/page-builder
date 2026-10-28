"""로컬 무API 회귀 테스트: 상품 정보 반영, 문법, 원고 품질 게이트, 출력 형식."""
import ast
import io
import json
import re
import unittest
import sys
from pathlib import Path
from types import SimpleNamespace

HERE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HERE))
from quality_engine import (
    TIP_PROFILES, check_quality, normalize_generated_strict,
    natural_wrap, check_structure, editorial_issues, score_editorial,
    build_editorial_brief, soften_unverified_claims, apply_grounded_safety, validate_required_inputs,
)


def app_functions():
    source = (HERE / 'app.py').read_text(encoding='utf-8')
    source_ast = ast.parse(source)
    keep = {'build_generation_prompt', 'generate_structured_copy', 'render_subsc_html',
            'render_text_source', 'build_subtap_html', 'assemble_final_output',
            'result_to_docx_bytes', 'add_blue_divider', 'extract_visual_evidence', 'format_material_line',
            'format_measurement_lines', 'extract_json', 'clean_line'}
    code = ast.Module(body=[n for n in source_ast.body if isinstance(n, ast.FunctionDef) and n.name in keep], type_ignores=[])
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from quality_engine import fact_report
    namespace = dict(json=json, io=io, re=re, TIP_PROFILES=TIP_PROFILES,
                     fact_report=fact_report, check_quality=check_quality, check_structure=check_structure,
                     build_editorial_brief=build_editorial_brief, editorial_issues=editorial_issues,
                     score_editorial=score_editorial, soften_unverified_claims=soften_unverified_claims,
                     apply_grounded_safety=apply_grounded_safety, validate_required_inputs=validate_required_inputs,
                     normalize_generated_strict=normalize_generated_strict,
                     natural_wrap=natural_wrap, Document=Document, OxmlElement=OxmlElement, qn=qn, Pt=Pt,
                     RateLimitError=RuntimeError, DIVIDER_MARKER="────────────────────────────────────────",
                     FIXED_HTML_HEAD='<meta charset="UTF-8">')
    exec(compile(code, str(HERE / 'app.py'), 'exec'), namespace)
    return namespace


def product(category='블라우스'):
    if category == '블라우스':
        return dict(display_name='비비안 브이넥 타이 블라우스 (3 color)', color='아이보리 / 연그레이 / 블랙',
                    size='FREE 77까지', material='텐셀 45% 레이온 37%',
                    material_desc_raw='부드러운 터치감, 은은한 광택', detail_tip='탈부착 타이, 소매 핀턱',
                    fit='팔뚝 군살 커버, 복부 여유핏', appeal_points='하객룩, 구김 적음', etc='',
                    target='4050 여성', washing='드라이클리닝 권장, 건조기 사용 금지',
                    measurement_lines=['어깨단면37 / 가슴둘레106 / 총장63.5'], country='한국')
    if category == '슬랙스':
        return dict(display_name='매그 S 639 부츠컷 슬랙스', color='네이비 / 블랙', size='S M L',
                    material='폴리에스터 95%, 스판 5%', material_desc_raw='텐션이 높고 복원력 좋음',
                    detail_tip='부츠컷, 허리 밴딩', fit='허벅지 여유, 다리선 슬림',
                    appeal_points='오피스룩, 구김 적음', etc='', target='4050 여성',
                    washing='찬물 손세탁 권장', measurement_lines=['S 허리둘레 66 / 총장 96'], country='한국')
    return dict(display_name='하이넥 스트링 야상', color='카키', size='FREE 77까지', material='코튼 100%',
                material_desc_raw='피치 가공 표면', detail_tip='하이넥, 밑단 스트링',
                fit='여유로운 어깨, 힙 덮는 기장', appeal_points='간절기 아우터, 여행', etc='',
                target='4050 여성', washing='단독 세탁', measurement_lines=['어깨단면 55 / 총장 70'], country='한국')


def draft_for(data):
    source = ' '.join(str(data[k]) for k in ('detail_tip', 'fit', 'appeal_points', 'material', 'material_desc_raw'))
    md_templates = {
      "choice": [
        f"{data['display_name']}의 가장 중요한 특징은 {data['detail_tip']}이라는 디테일입니다. 이런 구성이 필요한 장면을 명확하게 설명할 수 있습니다.",
        f"{data['appeal_points']}이라는 고객의 불편을 기준으로 해당 상품의 선택 이유를 구체적으로 살펴볼 수 있습니다. 단순한 광고 문구를 피합니다.",
        "실측과 상세 소재 설명을 먼저 확인하고, 평소 선호하는 옷과 비교하여 자신에게 적합한지를 판단하는 편이 좋습니다. 상품의 실제 정보가 선택 기준입니다."
      ],
      "fabric": [
        f"소재 구성은 {data['material']}이며 원단 설명에는 {data['material_desc_raw']}이라고 나와 있습니다. 이 정보를 바탕으로 감촉과 표면 특성을 이해할 수 있습니다.",
        "두께와 질감은 착용 계절을 가늠하는 중요한 기준이지만, 입력되지 않은 온도 적합성이나 기능을 확정적으로 안내하지 않습니다.",
        "관리 안내도 확인된 세탁법을 우선해야 하므로, 고객이 케어라벨 정보를 확인하도록 설명하는 것이 중요합니다."
      ],
      "fit": [
        f"핏에서 확인할 사항은 {data['fit']}으로 정리됩니다. 취향과 체형에 따라 같은 사이즈라도 느낌이 달라질 수 있습니다.",
        f"사이즈는 {data['size']}입니다. 평소 자주 입는 상품의 실측과 비교하면 원하는 실루엣을 이해하는 데 도움이 됩니다.",
        "착용 시 답답함과 움직임의 여유는 개인차가 있으므로, 정확한 사이즈 선택은 표기된 치수와 체형을 함께 고려해야 합니다."
      ],
      "occasion": [
        f"활용 장면은 {data['appeal_points']}을 참고하는 것이 좋습니다. 실제 디테일인 {data['detail_tip']}에 따라 코디 분위기가 달라집니다.",
        "평소 즐겨 입는 하의와 아우터를 함께 고려하면 활용도를 판단할 수 있고, 특별한 날에만 입게 되는 옷인지도 구분할 수 있습니다.",
        "자주 입는 색상과 계절감을 비교하여 조합을 선택하면 여러 상황에 맞춰 활용할 수 있으므로, 자신의 옷장 속 기본 아이템을 떠올려 보는 것이 좋습니다."
      ]
    }
    return {'material_desc_lines': [f'{source}로 이해할 수 있습니다.', f'{source}를 참고해 선택합니다.', f'{source}가 상품의 주요 특징입니다.'],
      'recommend_lines': [f'{data["detail_tip"]}을 찾는 분', f'{data["fit"]}을 원하는 분', '착용할 때의 편안함을 원하는 분', '관리의 편리함을 중요하게 생각하시는 분'],
      'review_lines': ['실측과 소재를 비교해 보기 좋은 상품이에요.', '평소 착용하는 옷과 비교하여 선택할 수 있어요.', '소재 정보와 핏을 함께 확인해보세요.', '상세 사이즈표가 선택에 도움이 됩니다.'],
      'faqs': [{'q': 'Q. 소재가 무엇인가요?', 'a': 'A. ' + str(data['material'])}, {'q': 'Q. 디테일은 어떤가요?', 'a': 'A. ' + str(data['detail_tip'])}, {'q': 'Q. 사이즈는 어떤가요?', 'a': 'A. ' + str(data['size'])}, {'q': 'Q. 세탁법은요?', 'a': 'A. ' + str(data['washing'])}],
      'md_sections': md_templates,
      'size_tips': {x: [f'{TIP_PROFILES[x]} 기준에서 실측과 기존 옷을 비교해 주세요.', f'{TIP_PROFILES[x]}은 개인차를 고려해 선택하세요.'] for x in TIP_PROFILES}}


class RepoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.f = app_functions()

    def test_unchanged_ui_and_downloads(self):
        s = (HERE/'app.py').read_text(encoding='utf-8')
        for piece in ['st.title("MISHARP PAGE BUILDER")', 'TXT 다운로드', 'HWP 다운로드',
                      'generated_result', 'generated_docx', 'DIVIDER_MARKER', 'ㅇ55 (90) 160cm 50kg']:
            self.assertIn(piece, s)
        self.assertNotIn('✓쇼핑에 꼭 참고하세요</h3>', s)

    def test_prompt_has_specific_product_facts(self):
        for category in ['블라우스', '슬랙스', '야상']:
            d=product(category)
            prompt=self.f['build_generation_prompt'](d, '사실 근거로 설명')
            for field in ('display_name','detail_tip','fit','appeal_points','material'):
                self.assertIn(str(d[field]), prompt)
            self.assertIn('입력된 한 상품', prompt)

    def test_generic_fallback_disallowed(self):
        d=product('블라우스')
        draft=draft_for(d)
        draft['recommend_lines'] = draft['recommend_lines'][:1]
        issues=check_quality(draft,d)
        self.assertTrue(any('recommend_lines' in s for s in issues))
        with self.assertRaises(ValueError):
            normalize_generated_strict(draft)

    def test_no_cross_product_text(self):
        for category in ['블라우스','슬랙스','야상']:
            d=product(category)
            result=normalize_generated_strict(draft_for(d))
            assembled=self.f['assemble_final_output'](d,result)
            self.assertIn(d['display_name'],assembled)
            for term in ('비비안', '매그', '야상'):
                if term not in d['display_name']:
                    self.assertNotIn(term, assembled)
            self.assertEqual(assembled.count('────────────────────────────────────────'), 1)
            self.assertNotIn('✓쇼핑에 꼭 참고하세요', assembled)
            self.assertNotIn('[구매 전 꼭 확인해 주세요]', assembled)
            for label in ('ㅇ55 (90) 160cm 50kg', 'ㅇ66 (95) 162cm 55kg', 'ㅇ66반 (95) 163cm 59kg','ㅇ77 (100) 161cm 62kg'):
                self.assertIn(label, assembled)

    def test_hwp_blue_line_and_txt_line(self):
        d=product()
        result=normalize_generated_strict(draft_for(d))
        assembled=self.f['assemble_final_output'](d,result)
        b=self.f['result_to_docx_bytes'](assembled)
        self.assertGreater(len(b),1000)
        # HWP 다운로드는 현재 레포가 내부적으로 생성하는 Word 호환 바이너리임 (기존 동작 유지)
        from docx import Document
        doc=Document(io.BytesIO(b))
        borders=[]
        for p in doc.paragraphs:
            pPr=p._p.pPr
            if pPr is not None and pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr') is not None:
                borders.append(p)
        self.assertEqual(len(borders),1)
        line=borders[0]._p.xml
        self.assertIn('2F6DF6', line)
        pos=next(i for i,p in enumerate(doc.paragraphs) if p._p.xml == borders[0]._p.xml)
        self.assertEqual([p.text for p in doc.paragraphs[pos-2:pos]],['',''])
        self.assertEqual([p.text for p in doc.paragraphs[pos+1:pos+3]],['',''])

    def test_content_gate_and_wrapping(self):
        d=product()
        draft=draft_for(d)
        self.assertFalse(check_quality(draft,d), check_quality(draft,d))
        self.assertTrue(all(len(x)<=24 for x in natural_wrap('이 원단은 은은하게 흐르는 질감이 자연스럽게 표현되어 체형에 따라 서로 다른 착용감이 느껴질 수 있습니다.',24)))

    def test_engine_draft_and_revision_integration_mock(self):
        d=product()
        initial=draft_for(d)
        incomplete=dict(initial, recommend_lines=initial['recommend_lines'][:1])
        calls=[]
        def fake_chat_with_retry(**kwargs):
            calls.append(kwargs)
            contents = json.dumps(incomplete if len(calls)==1 else initial, ensure_ascii=False)
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=contents))])
        self.f['chat_with_retry']=fake_chat_with_retry
        self.f['file_to_content_item']=lambda image: None
        result=self.f['generate_structured_copy'](d,'',[])
        self.assertLessEqual(len(calls),3)
        self.assertGreaterEqual(len(calls),2)
        self.assertEqual(len(result['recommend_lines']),4)
        self.assertIn('탈부착', str(calls[0]['messages']))
        self.assertIn('recommend_lines', str(calls[1]['messages']))


class QualityGateRegressionTests(unittest.TestCase):
    """10/28 뽀글이 재킷 오류·비정상 검수 결과·입력란 공란 재현."""
    @classmethod
    def setUpClass(cls):
        cls.f = app_functions()

    def jacket_data(self):
        return dict(
            display_name='보송 뽀글 웜니트 자켓 (2 color)', color='베이지, 차콜',
            size='FREE 사이즈로 77까지 추천드립니다.', material='폴리에스터100 (건조기 사용 금지)',
            material_desc_raw='촘촘하게 결이 살아 있는 폴리에스터 100% 뽀글이 원단, 도톰한 볼륨감, 포근한 질감, 털 빠짐이 거의 없습니다.',
            country='중국', measurement_lines=['어깨단면51 / 가슴둘레116 / 암홀둘레52 / 소매길이53 / 소매단둘레30 / 총길이58'],
            detail_tip='', fit='', appeal_points='', etc='', target='4050 여성',
            washing='드라이클리닝, 단독 울세탁, 손세탁 권장. 건조기 사용 금지',
        )

    def short_fit_draft(self, d):
        draft = draft_for(product('블라우스'))
        draft['md_sections']['fit'] = [
            'FREE 사이즈로 77까지 추천합니다.',
            '가슴둘레 116cm, 어깨단면 51cm로 표시되어 있습니다.',
            '실측을 비교하여 착용할 사이즈를 선택해 주세요.',
        ]
        # 본 테스트는 구조/분량 오류를 재현한다. 다른 상품 용어는 QA에서는 참고하지 않는다.
        return draft

    def test_short_fit_no_longer_blocks_result(self):
        d = self.jacket_data()
        draft = self.short_fit_draft(d)
        self.assertLess(sum(len(x) for x in draft['md_sections']['fit']),115)
        self.assertFalse(check_structure(draft), check_structure(draft))
        normalize_generated_strict(draft)
        # 사진, 선택한 원고가 있어도 115자 문턱에서 실패하지 않는다.
        responses = [draft, draft]
        calls = []
        def fake_chat_with_retry(**kwargs):
            calls.append(kwargs)
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(
                content=json.dumps(responses.pop(0), ensure_ascii=False)))])
        self.f['chat_with_retry'] = fake_chat_with_retry
        self.f['file_to_content_item'] = lambda image: {'type':'image_url','image_url':{'url':'data:image/jpeg;base64,ZmFrZQ=='}}
        generated = self.f['generate_structured_copy'](d, '', ['image1','image2'])
        self.assertLessEqual(len(calls),3)
        self.assertGreaterEqual(len(calls),2)
        self.assertEqual(generated['md_sections']['fit'], draft['md_sections']['fit'])
        self.assertEqual(len(calls[0]['messages'][1]['content']),3)
        self.assertEqual(len(calls[1]['messages'][1]['content']),3)
        output = self.f['assemble_final_output'](d,normalize_generated_strict(generated))
        self.assertIn('보송 뽀글', output)
        self.assertIn('가슴둘레 116cm',output)
        self.assertEqual(output.count('────────────────────────────────────────'),1)

    def test_editor_bad_json_keeps_original_if_usable(self):
        d = product('야상')
        initial = draft_for(d)
        responses = [json.dumps(initial,ensure_ascii=False), 'not a JSON reply']
        def fake_chat_with_retry(**kwargs):
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=responses.pop(0)))])
        self.f['chat_with_retry'] = fake_chat_with_retry
        self.assertEqual(self.f['generate_structured_copy'](d,'',[]),initial)

    def test_editor_service_error_keeps_usable_original(self):
        d=product('슬랙스')
        initial=draft_for(d)
        count=[0]
        def fake_chat_with_retry(**kwargs):
            count[0]+=1
            if count[0]==1:
                return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(initial,ensure_ascii=False)))])
            raise RuntimeError('simulated network failure')
        self.f['chat_with_retry']=fake_chat_with_retry
        self.assertEqual(self.f['generate_structured_copy'](d,'',[]),initial)

    def test_unusable_drafts_get_one_targeted_repair(self):
        d=product('블라우스')
        full=draft_for(d)
        broken=dict(full,recommend_lines=full['recommend_lines'][:1])
        responses=[broken, broken, full]
        calls=[]
        def fake_chat_with_retry(**kwargs):
            calls.append(kwargs)
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(responses.pop(0),ensure_ascii=False)))])
        self.f['chat_with_retry']=fake_chat_with_retry
        self.assertEqual(self.f['generate_structured_copy'](d,'',[]),full)
        self.assertEqual(len(calls),3)
        self.assertIn('추천',str(calls[2]['messages']))

    def test_editor_damaged_fields_falls_back_to_valid_initial(self):
        d=product('야상')
        full=draft_for(d)
        broken=dict(full,recommend_lines=full['recommend_lines'][:1])
        responses=[full,broken]
        def fake_chat_with_retry(**kwargs):
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(responses.pop(0),ensure_ascii=False)))])
        self.f['chat_with_retry']=fake_chat_with_retry
        self.assertEqual(self.f['generate_structured_copy'](d,'',[]),full)



class EditorialRegressionTests(unittest.TestCase):
    """Real 10/28 jacket outputs: misleading colors, exaggerated fur shedding and weak FAQs."""
    @staticmethod
    def jacket():
        return {
            'display_name':'보송 뽀글 립니트 자켓 (2 color)',
            'color':'베이지, 차콜','size':'FREE 77까지',
            'material':'폴리에스터100 (건조기 사용 금지)',
            'material_desc_raw':'양털 같은 포근한 뽀글이 원단, 부드러운 촉감, 털 빠짐이 거의 없습니다.',
            'detail_tip':'','fit':'','appeal_points':'', 'etc':'',
            'washing':'드라이클리닝, 손세탁 가능, 건조기 사용 금지',
            'measurement_lines':['어깨단면 51 / 가슴둘레 116 / 총길이 58'],
        }

    def test_warn_unknown_color_and_overstated_shedding(self):
        d=self.jacket()
        draft=draft_for(product('블라우스'))
        draft['recommend_lines']=['털 빠짐 없는 자켓을 찾으시는 분']*4
        draft['faqs']=[{'q':'Q. 소재가 무엇인가요?', 'a':'A. 폴리에스터'},
                       {'q':'Q. 사이즈가 어떻게 되나요?', 'a':'A. 77까지'},
                       {'q':'Q. 컬러 옵션은 무엇인가요?', 'a':'A. 아이보리 입니다'},
                       {'q':'Q. 건조기 가능하나요?', 'a':'A. 건조기는 금지입니다'}]
        findings=editorial_issues(draft,d)
        self.assertTrue(any('아이보리' in x for x in findings),findings)
        self.assertTrue(any('과장' in x for x in findings),findings)
        self.assertTrue(any('기본정보' in x for x in findings),findings)
        softened=soften_unverified_claims(draft,d)
        self.assertIn('털 빠짐이 거의 없는',softened['recommend_lines'][0])
        self.assertEqual(draft['recommend_lines'][0],'털 빠짐 없는 자켓을 찾으시는 분')

    def test_warn_unverified_review_experience(self):
        d=self.jacket()
        draft=draft_for(product('블라우스'))
        draft['review_lines'][0]='제가 직접 입어보니 너무 편해요.'
        issues=editorial_issues(draft,d)
        self.assertTrue(any('1인칭' in x for x in issues),issues)

    def test_warn_hip_cover_for_short_length(self):
        d=self.jacket()
        draft=draft_for(product('야상'))
        draft['md_sections']['fit']=['총길이 58cm, 힙을 덮는 기장으로 편하게 입습니다.']
        issues=editorial_issues(draft,d)
        self.assertTrue(any('힙' in x for x in issues),issues)

    def test_editor_prefers_quality_over_extra_length(self):
        d=self.jacket()
        base=draft_for(product('야상'))
        base['md_sections']['choice']=['뽀글이 소재를 고를 때 털 빠짐 관리가 걱정되셨다면, 털 빠짐이 거의 없는 원단을 먼저 살펴보세요.']
        base['recommend_lines']=['포근한 원단을 찾으시는 분', '피부 촉감을 중시하시는 분', '두 컬러 중 골라 입고 싶으신 분', '간절기에 가볍게 걸칠 자켓을 찾으시는 분']
        base['review_lines']=['원단의 질감과 밀도를 확인해 주세요.']*4
        base['faqs']=[
            {'q':'Q. 뽀글이 소재 관리가 번거롭지 않을까요?', 'a':'A. 건조기 사용은 금지하고 드라이클리닝 등 표시 세탁법을 따라주세요.'},
            {'q':'Q. 도톰한 이너를 입어도 괜찮을까요?', 'a':'A. 실측 가슴둘레 116cm를 기존 아우터와 비교해보세요.'},
            {'q':'Q. 평소 77을 입는데 부해 보일까요?', 'a':'A. 실측과 선호하는 실루엣에 따라 다르니 기존 자켓을 비교해 주세요.'},
            {'q':'Q. 베이지와 차콜 중 어떤 코디가 좋을까요?', 'a':'A. 평소 즐겨 입는 팬츠 색과 비교해 선택하시면 좋습니다.'},
        ]
        weaker=__import__('copy').deepcopy(base)
        weaker['faqs']=[
            {'q':'Q. 소재가 무엇인가요?', 'a':'A. 폴리에스터 100%입니다.'},
            {'q':'Q. 사이즈가 어떻게 되나요?', 'a':'A. FREE 사이즈입니다.'},
            {'q':'Q. 컬러 옵션은 무엇인가요?', 'a':'A. 아이보리와 차콜입니다.'},
            {'q':'Q. 세탁은 어떻게 하나요?', 'a':'A. 손세탁합니다.'},
        ]
        weaker['md_sections']['choice'] += ['베이지, 차콜 컬러 FREE 사이즈 폴리에스터 100% 소재.' for _ in range(6)]
        self.assertLess(score_editorial(base,d),score_editorial(weaker,d))

    def test_output_checks_not_fake_testimonials(self):
        d=product('야상')
        result=normalize_generated_strict(draft_for(d))
        assembled=app_functions()['assemble_final_output'](d,result)
        self.assertIn('✓ 착용 전 체크 포인트',assembled)
        self.assertNotIn('미리 입어 본 착용후기',assembled)
        self.assertNotIn('✓쇼핑에 꼭 참고하세요',assembled)
        self.assertEqual(assembled.count('────────────────────────────────────────'),1)

class ActualOutputEditorialTests(unittest.TestCase):
    """2026-10-28 jacket outputs: flag concrete failures that passed prior QA."""
    @staticmethod
    def jacket():
        return EditorialRegressionTests.jacket()

    def test_purchase_motive_grounded_in_input(self):
        from quality_engine import purchase_motive
        d=self.jacket()
        self.assertIn('털 빠짐', purchase_motive(d))
        self.assertIn('구김', purchase_motive(product('슬랙스')))
        self.assertNotIn('털 빠짐', purchase_motive(product('야상')))

    def test_actual_weak_jacket_flags_issues(self):
        d=self.jacket()
        draft=draft_for(product('야상'))
        draft['md_sections']['choice']=[
            '보송 뽀글 립니트 자켓 (2 color)은 포근한 질감과 도톰한 볼륨감으로 쌀쌀한 날씨에 적합합니다.',
            '털 빠짐이 거의 없고, 부드러운 촉감으로 피부에 편안하게 닿아 착용감이 우수합니다.',
            'FREE 사이즈로 다양한 체형에 잘 맞아 외출복으로 추천드립니다.',
        ]
        draft['md_sections']['fabric']=[
            '폴리에스터 100%의 뽀글이 원단은 포근한 질감을 제공합니다.',
            '털 빠짐이 거의 없고 부드러운 촉감입니다.',
        ]
        draft['md_sections']['fit']=['FREE 사이즈로 77까지 추천드리며 다양한 체형에 적합한 핏을 제공합니다.',
                                     '어깨단면 51cm, 총길이 58cm로 적당한 길이감이 특징입니다.']
        draft['review_lines']=['털 빠짐이 거의 없어 깔끔하게 착용할 수 있습니다.',
                               '폴리에스터 소재로 관리가 용이합니다.',
                               'FREE 사이즈로 다양한 체형에 잘 맞습니다.',
                               '보온성이 뛰어나 쌀쌀한 날씨에 적합합니다.']
        draft['size_tips']['55']=['어깨와 가슴둘레가 넉넉하여 편안하게 착용 가능합니다.',
                                  '총길이 58cm로 힙을 살짝 덮어주는 길이감입니다.']
        findings=editorial_issues(draft,d)
        self.assertTrue(any('상품명으로 시작' in x for x in findings), findings)
        self.assertTrue(any('핵심 가슴둘레' in x for x in findings), findings)
        self.assertTrue(any('반복' in x for x in findings), findings)
        self.assertTrue(any('보온성' in x for x in findings), findings)
        self.assertTrue(any('힙' in x for x in findings), findings)
        self.assertTrue(any('체크 포인트' in x for x in findings), findings)
        self.assertTrue(any('적당' in x for x in findings), findings)

    def test_bulky_question_not_answered_by_generic_size_claim(self):
        d=self.jacket()
        draft=draft_for(product('슬랙스'))
        draft['faqs'][0]={'q':'Q. 이 자켓은 부해 보일까요?',
                           'a':'A. 도톰한 볼륨감이 있지만, 실측 사이즈를 참고하면 다양한 체형에 적합합니다.'}
        issues=editorial_issues(draft,d)
        self.assertTrue(any('FAQ 답변' in x for x in issues), issues)

    def test_unverified_layering_and_length_claims_rejected(self):
        d=self.jacket()
        draft=draft_for(product('야상'))
        draft['size_tips']['66half']=['이너를 입어도 불편함이 없습니다.', '보온성이 뛰어난 실루엣입니다.']
        issues=editorial_issues(draft,d)
        self.assertTrue(any('이너' in x for x in issues),issues)
        self.assertTrue(any('보온성' in x for x in issues),issues)

    def test_prompt_sets_clear_customer_value_vs_spec_dump(self):
        d=self.jacket()
        prompt=app_functions()['build_generation_prompt'](d,'')
        for x in ('고객 고민','가슴둘레','입력값','체크 포인트','선택 이유'):
            self.assertIn(x,prompt)

    def test_targeted_third_revision_is_chosen_only_if_better(self):
        import copy
        d=self.jacket()
        weak=draft_for(product('야상'))
        weak['md_sections']['choice']=[d['display_name']+'은 포근하고 도톰합니다.']
        weak['faqs']=[{'q':'Q. 소재가 무엇인가요?', 'a':'A. 폴리에스터 100%입니다.'}]*4
        strong=copy.deepcopy(weak)
        strong['md_sections']['choice']=['뽀글이 소재의 털 빠짐이 걱정되셨다면 이 원단을 눈여겨보세요.',
                                          '털 빠짐이 거의 없다고 입력된 소재로 관리 부담을 덜 수 있습니다.']
        strong['faqs']=[{'q':'Q. 부해 보이지 않을까요?', 'a':'A. 도톰한 뽀글이 질감이라 체형에 따라 볼륨감이 다르게 느껴질 수 있습니다. 가슴둘레 실측 116cm를 평소 옷과 비교해 보세요.'},
                        {'q':'Q. 두꺼운 이너를 입어도 될까요?', 'a':'A. 이너 착용 여유는 확정하기 어렵습니다. 평소 자켓의 가슴둘레 116cm와 비교해 확인하세요.'},
                        {'q':'Q. 관리가 번거로운가요?', 'a':'A. 건조기는 금지입니다. 드라이클리닝과 손세탁 권장 안내를 확인해 주세요.'},
                        {'q':'Q. 평소 슬랙스와 코디가 될까요?', 'a':'A. 차콜은 차분한 코디, 베이지는 부드러운 톤의 코디에 활용해 보세요.'}]
        strong['review_lines']=['가슴둘레 116cm를 기존 자켓과 비교해 보세요.','총장 58cm를 평소 아우터와 비교해 보세요.',
                                '뽀글이 원단의 표면감을 확인해 보세요.','건조기를 피하고 케어라벨을 살펴보세요.']
        strong['md_sections']['fit']=['가슴둘레 116cm와 어깨단면 51cm 실측을 평소 자켓과 비교해 보세요.']
        calls=[]
        responses=[weak,weak,strong]
        def fake_chat_with_retry(**kwargs):
            calls.append(kwargs)
            result=responses.pop(0)
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(result,ensure_ascii=False)))])
        f=app_functions()
        f['chat_with_retry']=fake_chat_with_retry
        result=f['generate_structured_copy'](d,'',[])
        self.assertEqual(len(calls),3)
        self.assertEqual(result['md_sections']['choice'],strong['md_sections']['choice'])


if __name__ == '__main__':
    unittest.main()
