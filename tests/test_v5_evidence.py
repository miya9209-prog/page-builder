"""10/28 actual-jacket corrections, photo-grounding, mandatory data regressions."""
import ast
import copy
import io
import json
import re
import sys
import unittest
from pathlib import Path
from types import SimpleNamespace

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from quality_engine import (
    validate_required_inputs, apply_grounded_safety, editorial_issues, fact_report,
    purchase_motive, build_editorial_brief, numeric_measurement_issues,
)
from test_quality import app_functions, product, draft_for


def jacket():
    return dict(display_name='보송 뽀글 립니트 자켓 (2 color)', color='베이지, 차콜',
                country='중국', size='FREE 사이즈로 77까지 추천드립니다.',
                material='폴리에스터100 (건조기 사용 금지)',
                material_desc_raw='뽀글이 원단의 양털처럼 도톰한 볼륨감, 부드러운 털 결, 털 빠짐이 거의 없음',
                measurement_lines=['어깨단면51 / 가슴둘레116 / 암홀둘레52 / 소매길이53 / 소매단둘레30 / 총길이58'],
                detail_tip='', fit='', appeal_points='', etc='', washing='드라이클리닝, 손세탁 권장. 건조기 사용 금지',
                target='4050 여성')


class ActualJacketRegressions(unittest.TestCase):
    def test_country_is_required_before_api(self):
        d=jacket(); d['country']=''
        self.assertEqual(validate_required_inputs(d), ['제조국'])
        self.assertEqual(validate_required_inputs(jacket()), [])

    def test_output_never_silently_leaves_country_blank(self):
        f=app_functions(); d=jacket(); d['country']=''
        sample=draft_for(product('블라우스'))
        s=f['assemble_final_output'](d,sample)
        self.assertIn('제조국 : 미입력 (확인 필요)', s)
        d['country']='중국'
        self.assertIn('제조국 : 중국', f['assemble_final_output'](d,sample))

    def test_safety_removes_unsupported_hip_claim_both_directions(self):
        d=jacket(); draft=draft_for(product('블라우스'))
        draft['size_tips']['55']=['총장 58cm로 힙을 살짝 덮어주는 기장감입니다.', '힙을 덮지 않는 총장으로 다양한 하의에 어울립니다.']
        fixed=apply_grounded_safety(draft,d)
        text=' '.join(fixed['size_tips']['55'])
        self.assertNotIn('힙을 덮지 않는',text)
        self.assertNotIn('힙을 살짝 덮어',text)
        self.assertIn('비교',text)
        self.assertEqual(draft['size_tips']['55'][0],'총장 58cm로 힙을 살짝 덮어주는 기장감입니다.')

    def test_safety_corrects_unproven_sizing_and_shedding(self):
        d=jacket(); draft=draft_for(product('블라우스'))
        draft['recommend_lines'][0]='털 빠짐 걱정 없는 자켓을 찾으시는 분'
        draft['size_tips']['66']=['총장 58cm로 다양한 체형에 잘 맞습니다.', '가슴둘레 116cm입니다.']
        fixed=apply_grounded_safety(draft,d)
        self.assertIn('털 빠짐이 거의 없는',fixed['recommend_lines'][0])
        self.assertNotIn('다양한 체형에 잘 맞',fixed['size_tips']['66'][0])

    def test_warning_when_unsupported_shedding_wording(self):
        d=jacket(); draft=draft_for(product('블라우스'))
        draft['recommend_lines'][0]='털 빠짐 걱정 없는 자켓을 찾으시는 분'
        issues=editorial_issues(draft,d)
        self.assertTrue(any('걱정 없는' in issue for issue in issues),issues)

    def test_photo_fact_prioritized_without_fake_fit(self):
        d=jacket(); d['visual_details']=['소매 끝에 립니트 시보리', '앞면 중앙 지퍼 여밈']
        text=build_editorial_brief(d)
        self.assertIn('소매 끝에 립니트 시보리',text)
        self.assertIn('사진에서 관찰된 디자인 포인트',purchase_motive(d))
        self.assertIn('소매 끝에 립니트 시보리',fact_report(d))

    def test_vision_failure_returns_empty_not_fake_feature(self):
        fn=app_functions()['extract_visual_evidence']
        self.assertEqual(fn(jacket(),[]),[])
        self.assertEqual(fn(jacket(),['some fake image']),[])

    def test_vision_photo_json_and_filter(self):
        # Inject mock response into extracted app function global scope.
        f=app_functions(); fn=f['extract_visual_evidence']
        def stub_chat(**kw):
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({
                'visible_details':['소매 끝 립니트 조직','앞 지퍼 여밈','보온성이 뛰어남','소매 끝 립니트 조직']
            },ensure_ascii=False)))])
        fn.__globals__['chat_with_retry']=stub_chat
        result=fn(jacket(),[{'type':'image_url','image_url':{'url':'data:image/jpeg;base64,QQ=='}}])
        self.assertEqual(result,['소매 끝 립니트 조직','앞 지퍼 여밈'])

    def test_warn_wrong_numeric_measurements(self):
        d=jacket(); draft=draft_for(product('블라우스'))
        draft['md_sections']['fit'] = ['가슴둘레 118cm로 여유로운 품입니다.']
        issues=numeric_measurement_issues(draft,d)
        self.assertTrue(any('가슴둘레' in x and '118' in x for x in issues),issues)
        draft['md_sections']['fit'] = ['가슴둘레 116cm입니다.']
        self.assertFalse(any('가슴둘레' in x for x in numeric_measurement_issues(draft,d)))

    def test_warn_photo_fact_missing_from_copy(self):
        d=jacket(); d['visual_details']=['소매 끝에 립니트 시보리']
        draft=draft_for(product('블라우스'))
        # Draft mentions no distinctive photo evidence
        from quality_engine import check_quality
        issues=check_quality(draft,d)
        self.assertTrue(any('사진에서 확인된 외관' in x for x in issues),issues)

    def test_previous_copy_no_coloration_or_format_regression(self):
        f=app_functions(); s=f['render_text_source'](draft_for(product('블라우스')))
        self.assertEqual(s.count('────────────────────────────────────────'),1)
        self.assertNotIn('쇼핑에 꼭 참고하세요',s)
        self.assertIn('✓ (FAQ)',s)
        full=f['assemble_final_output'](jacket(),draft_for(product('블라우스')))
        self.assertEqual(full.count('ㅇ66반 (95) 163cm 59kg'),1)
        b=f['result_to_docx_bytes'](full)
        from docx import Document
        doc=Document(io.BytesIO(b))
        blues=[p for p in doc.paragraphs if p._p.xml.count('2F6DF6')]
        self.assertEqual(len(blues),1)
        self.assertIn('제조국 : 중국', '\n'.join(p.text for p in doc.paragraphs))


if __name__=='__main__': unittest.main()
